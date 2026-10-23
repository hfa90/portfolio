"""Exportação da lista de leads para CSV (Excel BR) e XLSX formatado."""
from __future__ import annotations

import csv
import io
import re

from .database import ROTULO_STATUS

COLUNAS = [
    ("nome", "Empresa", 34), ("categoria", "Categoria", 22), ("telefone", "Telefone", 17),
    ("whatsapp", "WhatsApp", 30), ("email", "E-mail", 28), ("site", "Site", 30),
    ("instagram", "Instagram", 30), ("facebook", "Facebook", 30), ("linkedin", "LinkedIn", 30),
    ("endereco", "Endereço completo", 50), ("rua", "Rua", 30), ("numero", "Número", 9),
    ("complemento", "Complemento", 16), ("bairro", "Bairro", 20), ("cidade", "Cidade", 18), ("uf", "UF", 5),
    ("cep", "CEP", 11), ("latitude", "Latitude", 11), ("longitude", "Longitude", 11),
    ("avaliacao", "Avaliação", 10), ("num_avaliacoes", "Nº avaliações", 13), ("horario", "Horário", 40),
    ("situacao", "Situação", 14), ("status", "Status (funil)", 15), ("score", "Score", 8),
    ("tags", "Tags", 18), ("notas", "Notas", 30), ("ultimo_contato", "Último contato", 18),
    ("maps_url", "Google Maps", 30), ("criado_em", "Capturado em", 18),
]


def _linha(lead: dict) -> dict:
    l = dict(lead)
    e164 = re.sub(r"\D", "", l.get("telefone_e164") or "")
    l["whatsapp"] = f"https://wa.me/{e164}" if e164 and l.get("celular") else ""
    l["status"] = ROTULO_STATUS.get(l.get("status"), l.get("status"))
    return l


def gerar_csv(leads: list[dict]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", quoting=csv.QUOTE_MINIMAL)
    w.writerow([c[1] for c in COLUNAS])
    for lead in leads:
        l = _linha(lead)
        valores = []
        for k, _, _ in COLUNAS:
            v = l.get(k)
            if v is None:
                v = ""
            elif k in ("avaliacao", "latitude", "longitude"):
                v = str(v).replace(".", ",")  # Excel em português usa vírgula decimal
            valores.append(v)
        w.writerow(valores)
    return buf.getvalue().encode("utf-8-sig")


def gerar_xlsx(leads: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    cab_fill = PatternFill("solid", fgColor="1E3A8A")
    cab_font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
    borda = Border(bottom=Side(style="thin", color="E4E7EC"))

    ws.append([c[1] for c in COLUNAS])
    for i, (_, _, larg) in enumerate(COLUNAS, 1):
        cel = ws.cell(row=1, column=i)
        cel.fill, cel.font = cab_fill, cab_font
        cel.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = larg
    ws.row_dimensions[1].height = 24

    links = {"whatsapp", "site", "instagram", "facebook", "linkedin", "maps_url"}
    for r, lead in enumerate(leads, 2):
        l = _linha(lead)
        for c, (k, _, _) in enumerate(COLUNAS, 1):
            v = l.get(k)
            cel = ws.cell(row=r, column=c, value=v if v not in (None, "") else None)
            cel.border = borda
            if k in links and v:
                cel.hyperlink = v
                cel.font = Font(color="1D4ED8", underline="single")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUNAS))}{max(1, len(leads) + 1)}"

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
